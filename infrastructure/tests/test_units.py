import socket
from datetime import datetime, timedelta
from uuid import uuid4

import fitz
import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook
from pydantic import ValidationError

from agent_common.contracts import Answers, InsightRequest
from agent_common.core import digest, monday_window, settings
from backend_service.documents import chunks, extract
from agent_common.embeddings import vector_literal
from tool_gateway_service.api import app
from tool_gateway_service.tools import fetch_public, public_addresses


def test_week_window_is_monday_eight_not_midnight():
    at = datetime.fromisoformat("2026-10-26T08:00:00+03:00")
    start, end = monday_window(at)
    assert start.isoformat() == "2026-10-19T08:00:00+03:00"
    assert end == at
    assert end - start == timedelta(days=7)
    assert monday_window(at - timedelta(seconds=1))[1] == start
    assert monday_window(at - timedelta(hours=2), upcoming=True)[1] == end
    assert monday_window(at + timedelta(seconds=1), upcoming=True)[1] == end + timedelta(days=7)


def test_window_requires_timezone():
    with pytest.raises(ValueError):
        monday_window(datetime(2026, 10, 26))


def test_chunks_bound_and_preserve_pages():
    text = "Деятельность компании. " * 300
    result = list(chunks([(text, {"page": 3})]))
    assert len(result) > 1
    assert all(len(content) <= 1200 and meta["page"] == 3 for _, content, meta in result)
    assert result[0][2]["end"] > result[1][2]["start"]
    assert [n for n, *_ in result] == list(range(len(result)))
    assert not list(chunks([("", {})]))


def test_docx_extraction(tmp_path):
    path = tmp_path / "company.docx"
    doc = Document()
    doc.add_paragraph("Производство оборудования")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Показатель", "Значение"
    table.cell(1, 0).text, table.cell(1, 1).text = "Сотрудники", "100"
    doc.save(path)
    sections = extract(path, path.name)
    assert "Производство" in sections[0][0]
    assert "Сотрудники" in sections[-1][0]
    assert sections[-1][1]["table"] == 1


def test_xlsx_preserves_header_and_row(tmp_path):
    path = tmp_path / "opaque-storage-key"  # Actual upload storage has no extension.
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Продажи"
    sheet.append(["Товар", "Количество, шт."])
    sheet.append(["Оборудование", 12])
    workbook.save(path)
    sections = extract(path, "company.xlsx")
    assert sections[0][1]["sheet"] == "Продажи"
    assert sections[0][1]["row"] == 2
    assert "Количество, шт.: 12" in sections[0][0]


def test_pdf_text_extraction(tmp_path):
    path = tmp_path / "test.pdf"
    with fitz.open() as pdf:
        page = pdf.new_page()
        page.insert_text((40, 40), "Manufacturing company. Equipment sales and services. " * 3)
        pdf.save(path)
    sections = extract(path, path.name)
    assert sections[0][1] == {"page": 1, "extraction": "text"}


def test_scanned_pdf_uses_ocr(tmp_path, monkeypatch):
    path = tmp_path / "scan.pdf"
    with fitz.open() as pdf:
        pdf.new_page()
        pdf.save(path)
    monkeypatch.setattr("backend_service.documents.ocr", lambda image: "Сканированный документ компании")
    result = extract(path, path.name)
    assert result[0][1]["extraction"] == "ocr"


def test_vector_validation():
    assert vector_literal([1.0] + [0.0] * 1023).startswith("[1.0,")
    for value in ([0.0], [0.0] * 1024, [2.0] + [0.0] * 1023, [float("nan")] * 1024, [float("inf")] * 1024):
        with pytest.raises(ValueError):
            vector_literal(value)


@pytest.mark.parametrize(
    "url",
    ["http://example.org", "https://user:pass@example.org", "https://example.org:8443", "file:///etc/passwd"],
)
def test_fetch_rejects_unsafe_scheme_or_port(url):
    with pytest.raises(ValueError):
        fetch_public(url)


@pytest.mark.parametrize("address", ["127.0.0.1", "10.0.0.1", "169.254.169.254", "::1", "192.168.1.2"])
def test_ssrf_rejects_non_public_dns(monkeypatch, address):
    monkeypatch.setattr(socket, "getaddrinfo", lambda *a, **kw: [(2, 1, 6, "", (address, 443))])
    with pytest.raises(ValueError):
        public_addresses("example.org")


def test_dns_rebinding_mixed_addresses_rejected(monkeypatch):
    monkeypatch.setattr(
        socket,
        "getaddrinfo",
        lambda *a, **kw: [(2, 1, 6, "", ("8.8.8.8", 443)), (2, 1, 6, "", ("127.0.0.1", 443))],
    )
    with pytest.raises(ValueError):
        public_addresses("example.org")


def test_gateway_requires_internal_auth(monkeypatch):
    monkeypatch.setattr(settings, "internal_token", "secret")
    response = TestClient(app).post("/tools/web.search", json={})
    assert response.status_code == 401


def test_body_limit_before_parsing():
    response = TestClient(app).post("/tools/web.search", content=b"x" * 140000)
    assert response.status_code == 413


def test_chunked_body_limit_without_content_length():
    response = TestClient(app).post(
        "/tools/web.search",
        content=iter([b"x" * 70000, b"x" * 70000]),
        headers={"Content-Type": "application/json"},
    )
    assert response.status_code == 413


def test_callback_contract_exclusive_inputs():
    request_id = uuid4()
    for args in ({}, {"ranking_run_id": uuid4(), "article_ids": ["a"]}):
        with pytest.raises(ValidationError):
            InsightRequest(request_id=request_id, **args)
    assert InsightRequest(request_id=request_id, article_ids=[]).article_ids == []


def test_short_onboarding():
    with pytest.raises(ValidationError):
        Answers(role="CFO", goals=["a", "b", "c", "d"])
    assert Answers(role="CFO").responsibilities == []


def test_digest_is_stable():
    assert digest({"b": 2, "a": 1}) == digest({"a": 1, "b": 2})
