"""Перевод файлов из новостей в текст: PDF, DOCX, XLSX, CSV/TXT.

Используется парсерами, когда ссылка новости ведёт прямо на файл (ЦБ:
«Вестник», указания) или когда к статье/посту приложены файлы.

Ограничения, чтобы один файл не раздувал базу и не тормозил сбор:
PDF — первые MAX_PDF_PAGES страниц, любой файл — до MAX_FILE_CHARS символов,
не больше MAX_FILES_PER_ITEM файлов на новость, файлы до MAX_FILE_BYTES.

PDF-сканы (нет текстового слоя или он «битый») распознаются tesseract, но
только быстро: первые OCR_MAX_PAGES страниц параллельно, и не дольше
OCR_BUDGET_SECONDS за весь запуск парсера — дальше сканы идут без текста.
Без установленного tesseract (например, на Windows вне Docker) OCR пропускается.
"""

import io
import logging
import os
import re
import shutil
import subprocess
import threading
import time
import zipfile
from concurrent.futures import ThreadPoolExecutor
from typing import Optional
from urllib.parse import unquote, urljoin, urlparse

import trafilatura

from common import http_get

MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_PDF_PAGES = 30
MAX_SHEET_ROWS = 500
MAX_FILE_CHARS = 50_000
MAX_FILES_PER_ITEM = 3

# Замер на указаниях ЦБ (Docker, 12 ядер): 72 DPI и 3 страницы параллельно —
# ~1,4 с на документ, текста столько же и такого же качества, как при 100 DPI
OCR_MAX_PAGES = 3          # у указаний и приказов на первых страницах — заголовок и суть
OCR_DPI = 72
OCR_PAGE_TIMEOUT = 10      # секунд на страницу
OCR_BUDGET_SECONDS = 30    # суммарно по всем сканам за запуск (~20 сканов), дальше — без OCR
_ocr_spent = 0.0
_ocr_lock = threading.Lock()

# pypdf пишет предупреждение на каждый нестандартный шрифт — в логах сбора это шум
logging.getLogger('pypdf').setLevel(logging.ERROR)

# ссылки, похожие на файл; /File/ — так отдаёт документы cbr.ru
FILE_LINK_RE = re.compile(r'\.(pdf|docx|xlsx|csv|txt)(\?|#|$)|/File/', re.I)
HTML_TYPES = ('text/html', 'application/xhtml', 'application/xml', 'text/xml',
              'application/rss', 'application/atom')


def is_file(response) -> bool:
    """Ответ сервера — файл, а не веб-страница."""
    content_type = response.headers.get('content-type', '').lower()
    return bool(content_type) and not content_type.startswith(HTML_TYPES)


def file_name(response) -> str:
    disposition = response.headers.get('content-disposition', '')
    match = re.search(r"filename\*=utf-8''([^;]+)", disposition, re.I) \
        or re.search(r'filename="?([^";]+)"?', disposition, re.I)
    if match:
        return unquote(match.group(1)).strip()
    return unquote(urlparse(response.url).path.rsplit('/', 1)[-1]) or 'file'


def response_to_text(response) -> Optional[str]:
    """Текст файла из уже скачанного ответа. None — не смогли прочитать."""
    if len(response.content) > MAX_FILE_BYTES:
        return None
    try:
        text = _extract(response.content, response.headers.get('content-type', ''),
                        file_name(response))
    except Exception:  # битый или зашифрованный файл — просто без текста
        return None
    if not text or _looks_garbled(text):
        return None
    if len(text) > MAX_FILE_CHARS:
        text = text[:MAX_FILE_CHARS] + '\n[…файл обрезан]'
    return text


def file_links(html: bytes, base_url: str) -> list[str]:
    """Ссылки на файлы в основном тексте страницы (меню и подвал не считаются)."""
    xml = trafilatura.extract(html, output_format='xml', include_links=True) or ''
    links = []
    for target in re.findall(r'target="([^"]+)"', xml):
        url = urljoin(base_url, target.replace('&amp;', '&'))
        if FILE_LINK_RE.search(url) and url not in links:
            links.append(url)
    return links


def files_to_text(urls: list[str]) -> str:
    """Скачивает файлы и склеивает их текст с подписями «[Файл: имя]»."""
    parts = []
    for url in urls[:MAX_FILES_PER_ITEM]:
        name, text = file_text(url)
        if text:
            parts.append(f'[Файл: {name}]\n{text}')
    return '\n\n'.join(parts)


def file_text(url: str) -> tuple[str, Optional[str]]:
    """(имя файла, текст) по ссылке; текст None, если файл не скачать или не прочитать."""
    # cbr.ru без www отвечает на файлы 403, с www — отдаёт
    url = re.sub(r'^(https?://)cbr\.ru/', r'\1www.cbr.ru/', url)
    response = http_get(url, attempts=1)  # приложение не главное — без повторов
    if response is None or not is_file(response):
        return '', None
    return file_name(response), response_to_text(response)


def _extract(content: bytes, content_type: str, name: str) -> str:
    content_type = content_type.lower()
    name = name.lower()
    if content[:4] == b'%PDF':
        return _pdf(content)
    if content[:2] == b'PK':  # docx и xlsx — это zip-архивы
        names = zipfile.ZipFile(io.BytesIO(content)).namelist()
        if 'word/document.xml' in names:
            return _docx(content)
        if 'xl/workbook.xml' in names:
            return _xlsx(content)
        return ''
    if content_type.startswith(('text/plain', 'text/csv')) or name.endswith(('.csv', '.txt')):
        for encoding in ('utf-8', 'cp1251'):
            try:
                return content.decode(encoding)
            except UnicodeDecodeError:
                continue
    return ''


def _pdf(content: bytes) -> str:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(content))
    pages = [page.extract_text() or '' for page in reader.pages[:MAX_PDF_PAGES]]
    text = '\n'.join(pages).strip()
    if not text or _looks_garbled(text):  # скан: текстового слоя нет или он «битый»
        return _ocr_pdf(content)
    if len(reader.pages) > MAX_PDF_PAGES:
        text += f'\n[…показаны {MAX_PDF_PAGES} из {len(reader.pages)} страниц]'
    return text


def _ocr_pdf(content: bytes) -> str:
    global _ocr_spent
    if not shutil.which('tesseract'):
        return ''
    with _ocr_lock:
        if _ocr_spent >= OCR_BUDGET_SECONDS:
            return ''
    start = time.monotonic()
    import pypdfium2 as pdfium
    pdf = pdfium.PdfDocument(content)
    total = len(pdf)
    count = min(total, OCR_MAX_PAGES)
    images = [_pgm(pdf[i].render(scale=OCR_DPI / 72, grayscale=True).to_numpy())
              for i in range(count)]
    pdf.close()
    with ThreadPoolExecutor(max_workers=count or 1) as pool:
        text = '\n'.join(pool.map(_tesseract, images)).strip()
    with _ocr_lock:
        _ocr_spent += time.monotonic() - start
    if text and total > count:
        text += f'\n[…распознаны {count} из {total} страниц скана]'
    return text


def _pgm(pixels) -> bytes:
    # серая картинка в простом формате PGM: tesseract читает его из stdin,
    # а Pillow для этого не нужен
    height, width = pixels.shape[:2]
    return f'P5\n{width} {height}\n255\n'.encode() + pixels.reshape(height, width).tobytes()


def _tesseract(image: bytes) -> str:
    try:
        result = subprocess.run(
            ['tesseract', 'stdin', 'stdout', '-l', 'rus'],
            input=image, capture_output=True, timeout=OCR_PAGE_TIMEOUT,
            env={**os.environ, 'OMP_THREAD_LIMIT': '1'})  # страницы и так идут параллельно
        return result.stdout.decode('utf-8', errors='replace')
    except subprocess.TimeoutExpired:
        return ''


def _docx(content: bytes) -> str:
    from docx import Document
    document = Document(io.BytesIO(content))
    lines = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            lines.append(' | '.join(cell.text.strip() for cell in row.cells))
    return '\n'.join(lines)


def _xlsx(content: bytes) -> str:
    from openpyxl import load_workbook
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    lines = []
    for sheet in workbook.worksheets:
        lines.append(f'# {sheet.title}')
        for row in sheet.iter_rows(max_row=MAX_SHEET_ROWS, values_only=True):
            cells = ['' if v is None else str(v) for v in row]
            if any(cells):
                lines.append(' | '.join(cells))
    return '\n'.join(lines)


def _looks_garbled(text: str) -> bool:
    # Сканы с плохим встроенным распознаванием дают «цЕIITPAЛЬHЬIЙ БAIIк»:
    # латиница и кириллица вперемешку внутри одного слова.
    words = re.findall(r'\w{4,}', text[:5000])
    if len(words) < 20:
        return False
    mixed = sum(bool(re.search('[a-zA-Z]', w)) and bool(re.search('[а-яА-ЯёЁ]', w))
                for w in words)
    return mixed / len(words) > 0.15
