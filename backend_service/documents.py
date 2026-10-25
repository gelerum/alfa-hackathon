"""Bounded extraction runs in a separate worker process, never in the API process."""

import io
import re
import zipfile
from pathlib import Path

import fitz
import pytesseract
from docx import Document
from openpyxl import load_workbook
from PIL import Image, ImageOps

from agent_common.core import db, digest, execute, j, one, settings
from agent_common.embeddings import config, embed, vector_literal

SUPPORTED = {".pdf", ".docx", ".xlsx", ".png", ".jpg", ".jpeg", ".tif", ".tiff", ".webp"}
EXTRACTION_VERSION = "structured-ocr-v1"
Image.MAX_IMAGE_PIXELS = 35_000_000


def check_archive(path):
    with zipfile.ZipFile(path) as archive:
        infos = archive.infolist()
        if len(infos) > 10000 or sum(i.file_size for i in infos) > 100_000_000:
            raise ValueError("Expanded archive exceeds limit")
        if any(i.flag_bits & 1 for i in infos):
            raise ValueError("Encrypted documents are not supported")


def ocr(image):
    if image.width * image.height > 35_000_000:
        raise ValueError("Image resolution exceeds OCR limit")
    image = ImageOps.exif_transpose(image).convert("RGB")
    return pytesseract.image_to_string(image, lang="rus+eng", timeout=90)


class BoundedSections(list):
    def __init__(self):
        super().__init__()
        self.characters = 0

    def append(self, item):
        self.characters += len(item[0])
        if self.characters > settings.max_extracted_chars:
            raise ValueError("Extracted text exceeds limit")
        super().append(item)


def extract(path: Path, filename: str):
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED:
        raise ValueError("Unsupported format")
    sections = BoundedSections()
    if ext == ".pdf":
        with fitz.open(path) as pdf:
            if pdf.needs_pass or len(pdf) > settings.max_pages:
                raise ValueError("Encrypted PDF or page limit exceeded")
            for number, page in enumerate(pdf):
                text = page.get_text()
                method = "text"
                if len(text.strip()) < 40:
                    if page.rect.width * page.rect.height * 9 > 35_000_000:
                        raise ValueError("PDF page resolution exceeds OCR limit")
                    pix = page.get_pixmap(matrix=fitz.Matrix(3, 3))  # 216 DPI for small report/table text
                    text = ocr(Image.open(io.BytesIO(pix.tobytes("png"))))
                    method = "ocr"
                sections.append((text, {"page": number + 1, "extraction": method}))
    elif ext == ".docx":
        check_archive(path)
        doc = Document(path)
        for n, paragraph in enumerate(doc.paragraphs):
            if paragraph.text.strip():
                sections.append((paragraph.text, {"paragraph": n + 1}))
        for n, table in enumerate(doc.tables):
            headers = [c.text for c in table.rows[0].cells] if table.rows else []
            for row_no, row in enumerate(table.rows[1:], 2):
                sections.append(
                    (
                        " | ".join(f"{h}: {c.text}" for h, c in zip(headers, row.cells)),
                        {"table": n + 1, "row": row_no},
                    )
                )
    elif ext == ".xlsx":
        check_archive(path)
        # Storage keys deliberately have no extension. A binary stream bypasses openpyxl's
        # filename-extension check, without trusting the user-supplied filename as a disk path.
        with path.open("rb") as source:
            workbook = load_workbook(source, read_only=True, data_only=True, keep_links=False)
            try:
                for sheet in workbook:
                    if (sheet.max_row or 0) > 20000 or (sheet.max_column or 0) > 100:
                        raise ValueError("Spreadsheet dimensions exceed limit")
                    sheet.reset_dimensions()  # XML dimensions can be absent or inaccurate.
                    rows = sheet.iter_rows(values_only=True)
                    headers = next(rows, ())
                    if len(headers) > 100:
                        raise ValueError("Spreadsheet column limit exceeded")
                    for n, row in enumerate(rows, 2):
                        if n > 20000 or len(row) > 100:
                            raise ValueError("Spreadsheet dimensions exceed limit")
                        if any(v is not None for v in row):
                            text = " | ".join(
                                f"{headers[i] if i < len(headers) and headers[i] is not None else i + 1}: {v}"
                                for i, v in enumerate(row)
                                if v is not None
                            )
                            sections.append(
                                (
                                    text,
                                    {"sheet": sheet.title, "row": n, "formula_policy": "cached-values-only"},
                                )
                            )
            finally:
                workbook.close()
    else:
        with Image.open(path) as image:
            frames = getattr(image, "n_frames", 1)
            if frames > settings.max_pages:
                raise ValueError("Image page limit exceeded")
            for frame in range(frames):
                image.seek(frame)
                sections.append((ocr(image.copy()), {"page": frame + 1, "extraction": "ocr"}))
    if sum(len(text) for text, _ in sections) > settings.max_extracted_chars:
        raise ValueError("Extracted text exceeds limit")
    if not any(text.strip() for text, _ in sections):
        raise ValueError("No usable text extracted")
    return sections


def chunks(sections, size=1200, overlap=150):
    """Conservative character ceiling for multilingual 2048-token embedding service.

    Do not silently truncate oversized input; split with overlap and preserve locators.
    """
    ordinal = 0
    for text, metadata in sections:
        text = re.sub(r"[ \t]+", " ", text.replace("\x00", "")).strip()
        start = 0
        while start < len(text):
            end = min(start + size, len(text))
            if end < len(text):
                boundary = text.rfind(" ", start + size // 2, end)
                if boundary > start:
                    end = boundary
            content = text[start:end].strip()
            if content:
                yield ordinal, content, {**metadata, "start": start, "end": end, "chunk_version": "1200c-v1"}
                ordinal += 1
            if end == len(text):
                break
            start = max(start + 1, end - overlap)


def index_chunks(
    *,
    sections,
    title,
    kind,
    company_id=None,
    document_id=None,
    article_id=None,
    content_version=None,
    published_at=None,
):
    with db() as conn:
        origin = (
            conn.execute(
                "SELECT canonical_url AS url FROM news WHERE article_id=%s", (article_id,)
            ).fetchone()
            if article_id
            else conn.execute(
                "SELECT source_url AS url FROM documents WHERE document_id=%s", (document_id,)
            ).fetchone()
        )
        source_url = origin["url"] if origin else None
        for ordinal, content, metadata in chunks(sections):
            existing = conn.execute(
                "SELECT chunk_id FROM rag_chunks WHERE ordinal=%s AND "
                "((document_id=%s) OR (article_id=%s AND content_version=%s))",
                (ordinal, document_id, article_id, content_version),
            ).fetchone()
            if existing:
                continue
            source_type = (
                "news" if kind == "news" else ("bank_analytics" if kind == "bank_analytics" else "document")
            )
            ref = {
                "ordinal": ordinal,
                "content_version": content_version,
                "locator": metadata,
                "document_id": document_id,
                "article_id": article_id,
            }
            source = conn.execute(
                "INSERT INTO evidence_sources(company_id,source_type,source_ref,excerpt,"
                "source_date,document_id,article_id,source_url) VALUES(%s,%s,%s,%s,%s,%s,%s,%s) RETURNING source_id",
                (company_id, source_type, j(ref), content, published_at, document_id, article_id, source_url),
            ).fetchone()
            conn.execute(
                "INSERT INTO rag_chunks(company_id,kind,document_id,article_id,content_version,ordinal,"
                "title,content,metadata,published_at,source_id) VALUES(%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)",
                (
                    company_id,
                    kind,
                    document_id,
                    article_id,
                    content_version,
                    ordinal,
                    title,
                    content,
                    j(metadata),
                    published_at,
                    source["source_id"],
                ),
            )


def process_document(payload):
    document = one("SELECT * FROM documents WHERE document_id=%s", (payload["document_id"],))
    if not document:
        raise ValueError("Document not found")
    if document["processing_status"] == "ready":
        return {"document_id": document["document_id"]}
    execute(
        "UPDATE documents SET processing_status='processing',processing_error=NULL WHERE document_id=%s",
        (document["document_id"],),
    )
    try:
        path = settings.storage_dir / document["storage_key"]
        if digest(path.read_bytes()) != document["content_sha256"]:
            raise ValueError("Stored document hash mismatch")
        sections = extract(path, document["original_filename"])
        index_chunks(
            sections=sections,
            title=document["title"],
            kind="bank_analytics" if document["source_type"] == "bank_analytics" else "company_document",
            company_id=document["company_id"],
            document_id=document["document_id"],
            published_at=document["published_at"],
        )
        execute(
            "UPDATE documents SET processing_status='ready',extraction_version=%s WHERE document_id=%s",
            (EXTRACTION_VERSION, document["document_id"]),
        )
    except Exception as exc:
        execute(
            "UPDATE documents SET processing_status='failed',processing_error=%s WHERE document_id=%s",
            (str(exc)[:1000], document["document_id"]),
        )
        raise
    return {"document_id": document["document_id"], "sections": len(sections)}


def index_news_batch():
    # A separate indexer leaves the ranker's article embeddings unchanged.
    with db() as conn:
        rows = conn.execute(
            "SELECT n.article_id,n.content_version,r.published_at,r.title,r.body,r.description "
            "FROM news n JOIN news_revisions r USING(article_id,content_version) "
            "WHERE n.duplicate_of_id IS NULL AND NOT EXISTS(SELECT 1 FROM rag_chunks c "
            "WHERE c.article_id=n.article_id AND c.content_version=n.content_version) "
            "ORDER BY n.fetched_at DESC LIMIT 25"
        ).fetchall()
    for row in rows:
        index_chunks(
            sections=[(row["body"] or row["description"] or row["title"], {})],
            title=row["title"],
            kind="news",
            article_id=row["article_id"],
            content_version=row["content_version"],
            published_at=row["published_at"],
        )


def embed_chunks_batch():
    cfg = config(chunk=True)
    with db() as conn:
        rows = conn.execute(
            "SELECT chunk_id,content FROM rag_chunks WHERE embedding IS NULL "
            "ORDER BY observed_at,ordinal LIMIT 4"
        ).fetchall()
    if not rows:
        return
    vectors = embed([r["content"] for r in rows], "document", cfg)
    with db() as conn:
        for row, vector in zip(rows, vectors):
            conn.execute(
                "UPDATE rag_chunks SET embedding=%s::vector,embedding_config_id=%s "
                "WHERE chunk_id=%s AND embedding IS NULL",
                (vector_literal(vector), cfg["embedding_config_id"], row["chunk_id"]),
            )
